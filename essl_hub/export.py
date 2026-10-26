"""Attendance over a date range for a set of people: Team's today column and
the three-sheet Excel and Daily Detail CSV exports (U8, U9).

One query loads every punch for the set and range; rows are grouped here and
each person-day goes through ``compute_day``, so the cost stays constant in
queries whatever the people x days product.
"""

from __future__ import annotations

import csv
import io
from collections import defaultdict
from dataclasses import dataclass
from datetime import date, datetime, timedelta

import openpyxl
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter

from .db import device_ids_of
from .hours import (
    CORRECTION_IGNORED, DEFAULT_DAY_START_HOUR, DayResult, Punch, compute_day, punch_from_row, work_day,
    work_day_window,
)

MAX_RANGE_DAYS = 93
# Spreadsheet apps treat these leading characters as a formula (CSV injection).
_FORMULA_PREFIXES = ("=", "+", "-", "@", "\t", "\r")

class RangeError(ValueError):
    """A from/to pair the export refuses; the message is shown to the user."""


def parse_range(from_raw: str | None, to_raw: str | None, default: date) -> tuple[date, date]:
    def parse(raw: str | None, label: str) -> date:
        if not raw:
            return default
        try:
            return date.fromisoformat(raw)
        except ValueError:
            raise RangeError(f"The {label} date is not a valid date. Use YYYY-MM-DD.") from None

    start, end = parse(from_raw, "From"), parse(to_raw, "To")
    if start > end:
        raise RangeError("The From date is after the To date. Swap them or pick a new range.")
    if (end - start).days + 1 > MAX_RANGE_DAYS:
        raise RangeError(
            f"That range is {(end - start).days + 1} days. Export at most {MAX_RANGE_DAYS} days at a time."
        )
    return start, end


def _load(
    database, people, start, end, *, target_minutes, today, now, day_start_hour
) -> tuple[dict, dict, dict]:
    """Results and sites seen per (person id, work day), plus the punches per
    (device user id, work day), from one punch query
    spanning the first day's window start to the last day's window end."""
    # Punches under any of a person's IDs group under their main ID.
    main_of = {uid: p["device_user_id"] for p in people for uid in device_ids_of(p)}
    rows = database.punches_for_device_users_range(
        list(main_of),
        work_day_window(start, day_start_hour)[0],
        work_day_window(end, day_start_hour)[1],
    )
    windows = database.device_windows()  # KTD3: once per report
    grouped: dict[tuple[str, date], list[Punch]] = defaultdict(list)
    seen_sites: dict[tuple[str, date], set[str]] = defaultdict(set)
    for row in rows:
        if row.get("correction") == CORRECTION_IGNORED:
            continue  # KTD4: HR ignored it; the raw list still shows it
        stamp = row["event_timestamp"]
        owner = main_of.get(row["device_user_id"], row["device_user_id"])
        key = (owner, work_day(stamp, day_start_hour))
        grouped[key].append(punch_from_row(row, windows))
        if row.get("site"):
            seen_sites[key].add(row["site"])
    results, sites = {}, {}
    for person in people:
        for day in _days(start, end):
            is_today = day == today
            key = (person["device_user_id"], day)
            results[(person["id"], day)] = compute_day(
                grouped.get(key, []),
                target_minutes=target_minutes,
                now=now if is_today else None,
                is_today=is_today,
            )
            sites[(person["id"], day)] = tuple(sorted(seen_sites.get(key, ())))
    return results, sites, grouped


def _days(start: date, end: date):
    day = start
    while day <= end:
        yield day
        day += timedelta(days=1)


def day_results(
    database,
    people: list[dict],
    start: date,
    end: date,
    *,
    target_minutes: int,
    today: date,
    now: datetime | None,
    day_start_hour: int = DEFAULT_DAY_START_HOUR,
) -> dict[tuple[int, date], DayResult]:
    """compute_day for every (person, work day) in the range, from one punch
    query. ``now`` is only used for ``today``, matching the single-day views."""
    return _load(
        database, people, start, end, target_minutes=target_minutes, today=today, now=now,
        day_start_hour=day_start_hour,
    )[0]


# --- The attendance report (U9) ------------------------------------------

DAILY_HEADER = (
    "Emp ID", "Name", "Manager", "Site(s)", "Date", "Weekday", "First In", "Last Out",
    "Punches", "Gross", "Break", "Effective", "Effective (h)", "Status",
    "Shortfall vs target", "Remarks",
)
MATRIX_TOTALS = ("Present", "Absent", "Review", "Incomplete", "Total effective", "Days below target")
SUMMARY_HEADER = (
    "Emp ID", "Name", "Manager", "Days present", "Average first in", "Average effective",
    "Days below target", "Incomplete days", "Review days", "Expected hours", "Actual hours",
)
_DURATION_FORMAT = "[h]:mm:ss"
_TIME_FORMAT = "hh:mm:ss"
_WEEKEND_FILL = PatternFill("solid", fgColor="EEEEEE")
_BOLD = Font(bold=True)


@dataclass(frozen=True)
class DayRow:
    """One person-day, already reduced to what every sheet shows."""

    person: dict
    day: date
    result: DayResult
    sites: tuple[str, ...]
    code: str  # P, A, WO, R, I, O, or "" before the person existed
    status: str
    first_in: datetime | None
    last_out: datetime | None
    punches: int
    # None where the number would be fabricated (incomplete, absent, ...).
    gross: timedelta | None
    brk: timedelta | None
    effective: timedelta | None
    shortfall: timedelta | None


@dataclass(frozen=True)
class Report:
    start: date
    end: date
    day_start_hour: int
    target_minutes: int
    people: list[dict]  # by name
    manager_names: dict[int, str]
    rows: list[DayRow]  # by date, then name

    @property
    def days(self) -> list[date]:
        return list(_days(self.start, self.end))

    @property
    def note(self) -> str:
        hour = f"{self.day_start_hour:02d}:00"
        return (
            f"Attendance {self.start.isoformat()} to {self.end.isoformat()}. "
            f"Each work day runs {hour} to {hour} the next day (IST)."
        )

    def manager(self, person: dict) -> str:
        return self.manager_names.get(person.get("manager_id"), "")


def _since(person: dict) -> date | None:
    created = person.get("created_at")
    if isinstance(created, datetime):
        return created.date()
    if isinstance(created, date):
        return created
    return None


def _incomplete_label(result: DayResult) -> str:
    exits = any(a.startswith("Exit at") for a in result.anomalies)
    entries = any(a.startswith("No exit") for a in result.anomalies)
    if entries and not exits:
        return "Missing exit"
    if exits and not entries:
        return "Missing entry"
    return "Incomplete"


def _day_row(
    person: dict, day: date, result: DayResult, sites: tuple[str, ...], punches: list[Punch], target_seconds: int
) -> DayRow:
    closed_ends = [s.end for s in result.segments if s.end is not None]
    # Start of the first paired or open entry segment, mirroring Last Out (end
    # of the last paired segment); a stray swipe before it is not First In.
    first_in = min((s.start for s in result.segments), default=None)
    last_out = max(closed_ends, default=None)
    gross = brk = effective = shortfall = None
    weekend = day.weekday() >= 5
    since = _since(person)
    if result.status == "empty":
        if since is not None and day < since:
            code, status = "", ""
        elif weekend:
            code, status = "WO", "Weekend off"
        else:
            code, status = "A", "Absent"
    elif result.status == "incomplete":
        code, status = "I", _incomplete_label(result)
    else:
        gross = timedelta(seconds=result.gross_seconds or 0)
        brk = timedelta(seconds=result.break_seconds)
        effective = timedelta(seconds=result.effective_seconds)
        if result.status == "open":
            code, status = "O", "Open"
            last_out = None
        else:
            shortfall = timedelta(seconds=max(0, target_seconds - result.effective_seconds))
            if result.anomalies:
                # Gross spans first to last pair, so an unmatched mid-day swipe
                # would read as break: leave both blank, keep paired Effective.
                gross = brk = None
                code = "R"
                status = ("Met" if result.status == "met" else "Short") + " · review"
            else:
                code, status = "P", "Met" if result.status == "met" else "Short"
    return DayRow(
        person=person, day=day, result=result, sites=sites, code=code, status=status,
        first_in=first_in, last_out=last_out, punches=len(punches),
        gross=gross, brk=brk, effective=effective, shortfall=shortfall,
    )


def build_report(
    database,
    people: list[dict],
    start: date,
    end: date,
    *,
    manager_names: dict[int, str],
    target_minutes: int,
    today: date,
    now: datetime | None,
    day_start_hour: int = DEFAULT_DAY_START_HOUR,
) -> Report:
    """Every sheet's rows from one punch query."""
    results, sites, grouped = _load(
        database, people, start, end, target_minutes=target_minutes, today=today, now=now,
        day_start_hour=day_start_hour,
    )
    ordered = sorted(people, key=lambda p: (p["name"].lower(), p["id"]))
    rows = [
        _day_row(
            person, day, results[(person["id"], day)], sites[(person["id"], day)],
            grouped.get((person["device_user_id"], day), []), target_minutes * 60,
        )
        for day in _days(start, end)
        for person in ordered
    ]
    return Report(
        start=start, end=end, day_start_hour=day_start_hour, target_minutes=target_minutes,
        people=ordered, manager_names=manager_names, rows=rows,
    )


def safe_text(value: object) -> object:
    """Neutralize text a spreadsheet would run as a formula. Numbers pass."""
    if isinstance(value, str) and value.startswith(_FORMULA_PREFIXES):
        return "'" + value
    return value


def _hours(value: timedelta | None) -> float | None:
    return None if value is None else round(value.total_seconds() / 3600, 2)


def _hm(value: timedelta | None) -> str:
    if value is None:
        return ""
    minutes = int(value.total_seconds()) // 60
    return f"{minutes // 60}:{minutes % 60:02d}"


def _daily_values(report: Report, row: DayRow) -> list[object]:
    """Daily Detail in DAILY_HEADER order; the times and durations are
    native values, which each writer formats."""
    return [
        row.person["device_user_id"], row.person["name"], report.manager(row.person),
        ", ".join(row.sites), row.day, f"{row.day:%a}", row.first_in, row.last_out,
        row.punches, row.gross, row.brk, row.effective, _hours(row.effective), row.status,
        row.shortfall, "; ".join(row.result.anomalies),
    ]


def to_csv(report: Report) -> str:
    """Daily Detail only, UTF-8 with a BOM so Excel opens it on double-click."""
    output = io.StringIO(newline="")
    writer = csv.writer(output)
    writer.writerow(DAILY_HEADER)
    for row in report.rows:
        cells = []
        for value in _daily_values(report, row):
            if value is None:
                cells.append("")
            elif isinstance(value, datetime):
                cells.append(f"{value:%H:%M}")
            elif isinstance(value, date):
                cells.append(value.isoformat())
            elif isinstance(value, timedelta):
                cells.append(_hm(value))
            else:
                cells.append(safe_text(value))
        writer.writerow(cells)
    return "﻿" + output.getvalue()


def _write(sheet, row_index: int, values: list[object], formats: dict[int, str] | None = None) -> None:
    """Write one row; text is always a string cell, never a formula."""
    for column, value in enumerate(values, start=1):
        cell = sheet.cell(row=row_index, column=column, value=safe_text(value))
        if isinstance(cell.value, str):
            cell.data_type = "s"
        elif isinstance(value, timedelta):
            cell.number_format = _DURATION_FORMAT
        if formats and column in formats and value is not None:
            cell.number_format = formats[column]


def _titled(sheet, report: Report, header: tuple[str, ...] | list[str], extra_note: str = "") -> None:
    """Row 1 carries the range and work-day window; the header is row 3."""
    _write(sheet, 1, [report.note + (" " + extra_note if extra_note else "")])
    sheet.cell(row=1, column=1).font = _BOLD
    _write(sheet, 3, list(header))
    for cell in sheet[3]:
        cell.font = _BOLD


def _daily_sheet(sheet, report: Report) -> None:
    sheet.title = "Daily Detail"
    _titled(sheet, report, DAILY_HEADER)
    formats = {5: "yyyy-mm-dd", 7: _TIME_FORMAT, 8: _TIME_FORMAT, 13: "0.00"}
    for i, row in enumerate(report.rows, start=4):
        _write(sheet, i, _daily_values(report, row), formats)
    sheet.freeze_panes = "C4"
    sheet.auto_filter.ref = f"A3:{get_column_letter(len(DAILY_HEADER))}{max(3, sheet.max_row)}"
    widths = (10, 24, 22, 14, 11, 9, 10, 10, 8, 10, 10, 10, 12, 13, 12, 36)
    for column, width in enumerate(widths, start=1):
        sheet.column_dimensions[get_column_letter(column)].width = width


def _cell_text(row: DayRow) -> str:
    if row.effective is None:
        return row.code
    seconds = int(row.effective.total_seconds())
    return f"{row.code} {seconds // 3600}:{seconds % 3600 // 60:02d}"


def _below_target(row: DayRow) -> bool:
    return row.code in ("P", "R") and row.result.status == "short"


def _matrix_sheet(sheet, report: Report) -> None:
    sheet.title = "Monthly Matrix"
    days = report.days
    header = ["Emp ID", "Name"] + [f"{d:%a %d %b}" for d in days] + list(MATRIX_TOTALS)
    _titled(sheet, report, header, "Codes: P present, A absent, WO weekend off, R review, I incomplete, O open.")
    by_person: dict[int, list[DayRow]] = defaultdict(list)
    for row in report.rows:
        by_person[row.person["id"]].append(row)
    for i, person in enumerate(report.people, start=4):
        rows = by_person[person["id"]]
        codes = [r.code for r in rows]
        total = sum((r.effective for r in rows if r.effective is not None and r.code != "O"), timedelta())
        values = [person["device_user_id"], person["name"]] + [_cell_text(r) for r in rows] + [
            sum(c in ("P", "R", "O") for c in codes), codes.count("A"), codes.count("R"),
            codes.count("I"), total, sum(_below_target(r) for r in rows),
        ]
        _write(sheet, i, values)
    for offset, day in enumerate(days, start=3):
        if day.weekday() >= 5:
            for row_index in range(3, sheet.max_row + 1):
                sheet.cell(row=row_index, column=offset).fill = _WEEKEND_FILL
        sheet.column_dimensions[get_column_letter(offset)].width = 11
    sheet.column_dimensions["B"].width = 24
    sheet.freeze_panes = "C4"


def _summary_sheet(sheet, report: Report) -> None:
    sheet.title = "Person Summary"
    _titled(
        sheet, report, SUMMARY_HEADER,
        "Average first in and average effective count complete days only (met, short or review). "
        "Expected hours count weekdays from the day the person was added.",
    )
    by_person: dict[int, list[DayRow]] = defaultdict(list)
    for row in report.rows:
        by_person[row.person["id"]].append(row)
    window_start = timedelta(hours=report.day_start_hour)
    for i, person in enumerate(report.people, start=4):
        rows = by_person[person["id"]]
        complete = [r for r in rows if r.code in ("P", "R")]
        # Average first in as an offset into the work day, so a 23:50 and a
        # 00:10 start average to midnight, not noon.
        offsets = [r.first_in - datetime.combine(r.day, datetime.min.time()) - window_start for r in complete]
        average_in = None
        if offsets:
            mean = sum(offsets, timedelta()) / len(offsets) + window_start
            average_in = (datetime.min + timedelta(seconds=int(mean.total_seconds()) % 86400)).time()
        average_effective = (
            timedelta(seconds=sum(int(r.effective.total_seconds()) for r in complete) // len(complete))
            if complete else None
        )
        expected_days = sum(1 for r in rows if r.code != "" and r.day.weekday() < 5)
        actual = sum((r.effective for r in rows if r.effective is not None and r.code != "O"), timedelta())
        values = [
            person["device_user_id"], person["name"], report.manager(person),
            sum(r.code in ("P", "R", "O") for r in rows), average_in, average_effective,
            sum(_below_target(r) for r in rows), sum(r.code == "I" for r in rows),
            sum(r.code == "R" for r in rows),
            round(expected_days * report.target_minutes / 60, 2), _hours(actual),
        ]
        _write(sheet, i, values, {5: _TIME_FORMAT, 10: "0.00", 11: "0.00"})
    for column, width in enumerate((10, 24, 22, 13, 16, 17, 17, 15, 12, 15, 13), start=1):
        sheet.column_dimensions[get_column_letter(column)].width = width
    sheet.freeze_panes = "C4"


def to_xlsx(report: Report) -> bytes:
    workbook = openpyxl.Workbook()
    _daily_sheet(workbook.active, report)
    _matrix_sheet(workbook.create_sheet(), report)
    _summary_sheet(workbook.create_sheet(), report)
    output = io.BytesIO()
    workbook.save(output)
    return output.getvalue()
