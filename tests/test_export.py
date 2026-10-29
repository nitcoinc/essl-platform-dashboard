from __future__ import annotations

import csv
import io
import os
import tempfile
import unittest
from datetime import date, datetime, timedelta
from unittest import mock

import openpyxl

from essl_hub import db as db_module
from essl_hub import export
from essl_hub.config import DEVICES


class ExportTest(unittest.TestCase):
    def setUp(self) -> None:
        tmp = tempfile.TemporaryDirectory()
        self.addCleanup(tmp.cleanup)
        self.database = db_module.Database(os.path.join(tmp.name, "essl.db"))
        self.database.initialize(DEVICES)
        self.day = date(2026, 9, 1)

    def _people(self, count: int) -> list[dict]:
        for i in range(count):
            self.database.upsert_employee(device_user_id=str(100 + i), name=f"Person {i:03d}", email=f"p{i}@x.com")
        return self.database.employees()

    def _punch(self, uid: str, at: datetime, direction: str, n: int) -> dict:
        return {
            "fingerprint": f"fp-{uid}-{n}", "device_id": DEVICES[0].id, "site": "s", "device_serial": "x",
            "device_user_id": uid, "employee_name": "", "event_timestamp": at, "raw_state": "0",
            "door": "1", "reader": "", "raw_direction_byte": 0, "trusted_direction": direction, "raw": {},
        }

    def _report(self, people, start, end, **kwargs):
        options = dict(manager_names={}, target_minutes=480, today=date(2026, 10, 10), now=None)
        options.update(kwargs)
        return export.build_report(self.database, people, start, end, **options)

    def _at(self, days=0, hours=0, minutes=0, seconds=0) -> datetime:
        return datetime.combine(self.day, datetime.min.time()) + timedelta(
            days=days, hours=hours, minutes=minutes, seconds=seconds
        )

    def _book(self, report):
        return openpyxl.load_workbook(io.BytesIO(export.to_xlsx(report)))

    @staticmethod
    def _col(name: str) -> int:
        return export.DAILY_HEADER.index(name)

    def test_overnight_shift_is_one_row_with_next_day_last_out(self) -> None:
        people = self._people(1)
        self.database.insert_punches([
            self._punch("100", self._at(hours=22, minutes=10), "entry", 1),
            self._punch("100", self._at(days=1, hours=3, minutes=40), "exit", 2),
        ])
        first, second = self._report(people, self.day, self.day + timedelta(days=1)).rows
        self.assertEqual(first.day, self.day)
        self.assertEqual(first.last_out, self._at(days=1, hours=3, minutes=40))
        self.assertEqual(first.effective, timedelta(hours=5, minutes=30))
        self.assertEqual((first.status, first.code, first.punches), ("Short", "P", 2))
        self.assertEqual(first.shortfall, timedelta(hours=2, minutes=30))
        self.assertEqual(second.punches, 0)

    def test_flagged_day_counts_paired_time_and_shows_review(self) -> None:
        people = self._people(1)
        self.database.insert_punches([
            self._punch("100", self._at(hours=9), "entry", 1),
            self._punch("100", self._at(hours=10), "entry", 2),
            self._punch("100", self._at(hours=18), "exit", 3),
        ])
        report = self._report(people, self.day, self.day)
        (row,) = report.rows
        self.assertEqual((row.status, row.code, row.punches), ("Met · review", "R", 3))
        self.assertEqual(row.effective, timedelta(hours=8))
        # First In is the first paired segment's start, not the stray 09:00.
        self.assertEqual(row.first_in, self._at(hours=10))
        self.assertIsNone(row.gross)
        self.assertIsNone(row.brk)
        self.assertEqual(row.shortfall, timedelta(0))
        sheet = self._book(report)["Daily Detail"]
        self.assertEqual(sheet.cell(row=4, column=self._col("Remarks") + 1).value, "No exit after 09:00")
        self.assertEqual(self._book(report)["Monthly Matrix"].cell(row=4, column=3).value, "R 8:00")

    def test_mid_day_unmatched_swipe_does_not_become_break(self) -> None:
        people = self._people(1)
        self.database.insert_punches([
            self._punch("100", self._at(hours=9), "entry", 1),
            self._punch("100", self._at(hours=12), "exit", 2),
            self._punch("100", self._at(hours=13), "exit", 3),
            self._punch("100", self._at(hours=14), "entry", 4),
            self._punch("100", self._at(hours=17), "exit", 5),
        ])
        report = self._report(people, self.day, self.day)
        (row,) = report.rows
        self.assertEqual((row.status, row.code), ("Short · review", "R"))
        self.assertEqual(row.effective, timedelta(hours=6))
        self.assertEqual(row.shortfall, timedelta(hours=2))
        self.assertEqual(row.first_in, self._at(hours=9))
        values = [c.value for c in self._book(report)["Daily Detail"][4]]
        self.assertIsNone(values[self._col("Gross")])
        self.assertIsNone(values[self._col("Break")])
        self.assertEqual(values[self._col("Status")], "Short · review")
        self.assertIn("13:00", values[self._col("Remarks")])
        self.assertEqual(self._book(report)["Monthly Matrix"].cell(row=4, column=3).value, "R 6:00")
        summary = [c.value for c in self._book(report)["Person Summary"][4]]
        self.assertEqual(summary[export.SUMMARY_HEADER.index("Review days")], 1)
        self.assertEqual(summary[export.SUMMARY_HEADER.index("Days below target")], 1)
        self.assertEqual(summary[export.SUMMARY_HEADER.index("Actual hours")], 6.0)

    def test_first_in_blank_without_entry_segment(self) -> None:
        people = self._people(1)
        self.database.insert_punches([self._punch("100", self._at(hours=18), "exit", 1)])
        (row,) = self._report(people, self.day, self.day).rows
        self.assertEqual(row.code, "I")
        self.assertIsNone(row.first_in)

    def test_incomplete_day_blanks_numbers_and_says_missing_exit(self) -> None:
        people = self._people(1)
        self.database.insert_punches([self._punch("100", self._at(hours=9), "entry", 1)])
        report = self._report(people, self.day, self.day)
        (row,) = report.rows
        self.assertEqual((row.status, row.code), ("Missing exit", "I"))
        self.assertEqual(row.first_in, self._at(hours=9))
        self.assertIsNone(row.last_out)
        self.assertIsNone(row.effective)
        self.assertIsNone(row.shortfall)
        values = [c.value for c in self._book(report)["Daily Detail"][4]]
        for name in ("Last Out", "Gross", "Break", "Effective", "Effective (h)", "Shortfall vs target"):
            self.assertIsNone(values[self._col(name)], name)
        self.assertEqual(self._book(report)["Monthly Matrix"].cell(row=4, column=3).value, "I")

    def test_open_today_has_no_shortfall(self) -> None:
        people = self._people(1)
        self.database.insert_punches([self._punch("100", self._at(hours=9), "entry", 1)])
        (row,) = self._report(people, self.day, self.day, today=self.day, now=self._at(hours=11)).rows
        self.assertEqual((row.status, row.code), ("Open", "O"))
        self.assertEqual(row.effective, timedelta(hours=2))
        self.assertIsNone(row.shortfall)
        self.assertIsNone(row.last_out)

    def test_absent_weekend_and_before_created_codes(self) -> None:
        (person,) = self._people(1)
        person["created_at"] = datetime(2026, 9, 3, 12, 0)
        report = self._report([person], self.day, date(2026, 9, 7))  # Tue..Mon
        self.assertEqual([r.code for r in report.rows], ["", "", "A", "A", "WO", "WO", "A"])
        self.assertEqual(report.rows[4].status, "Weekend off")
        matrix = self._book(report)["Monthly Matrix"]
        self.assertEqual([c.value for c in matrix[4]][2:9], [None, None, "A", "A", "WO", "WO", "A"])
        absent = 3 + len(report.days) + export.MATRIX_TOTALS.index("Absent")
        self.assertEqual(matrix.cell(row=4, column=absent).value, 3)
        self.assertEqual(matrix.cell(row=3, column=7).fill.fgColor.rgb[-6:], "EEEEEE")  # Sat shaded
        summary = self._book(report)["Person Summary"]
        expected = export.SUMMARY_HEADER.index("Expected hours") + 1
        self.assertEqual(summary.cell(row=4, column=expected).value, 24)  # 3 weekdays x 8h

    def test_xlsx_cell_types_keep_seconds(self) -> None:
        people = self._people(1)
        self.database.insert_punches([
            self._punch("100", self._at(hours=22, minutes=10), "entry", 1),
            self._punch("100", self._at(days=1, hours=3, minutes=40, seconds=15), "exit", 2),
        ])
        book = self._book(self._report(people, self.day, self.day))
        self.assertEqual(book.sheetnames, ["Daily Detail", "Monthly Matrix", "Person Summary"])
        sheet = book["Daily Detail"]
        self.assertIn("2026-09-01 to 2026-09-01", sheet["A1"].value)
        self.assertIn("06:00 to 06:00", sheet["A1"].value)
        self.assertEqual([c.value for c in sheet[3]], list(export.DAILY_HEADER))
        self.assertEqual(sheet.freeze_panes, "C4")
        self.assertEqual(sheet.auto_filter.ref, "A3:P4")
        first_in = sheet.cell(row=4, column=self._col("First In") + 1)
        self.assertEqual(first_in.value, self._at(hours=22, minutes=10))
        self.assertEqual(first_in.number_format, "hh:mm:ss")
        last_out = sheet.cell(row=4, column=self._col("Last Out") + 1)
        self.assertEqual(last_out.value, self._at(days=1, hours=3, minutes=40, seconds=15))
        effective = sheet.cell(row=4, column=self._col("Effective") + 1)
        self.assertEqual(effective.number_format, "[h]:mm:ss")
        self.assertEqual(effective.value, timedelta(hours=5, minutes=30, seconds=15))
        self.assertEqual(sheet.cell(row=4, column=self._col("Effective (h)") + 1).value, 5.5)
        summary = book["Person Summary"]
        self.assertEqual(summary.cell(row=4, column=5).value.strftime("%H:%M"), "22:10")

    def test_summary_average_first_in_crosses_midnight(self) -> None:
        people = self._people(1)
        self.database.insert_punches([
            self._punch("100", self._at(hours=23, minutes=50), "entry", 1),
            self._punch("100", self._at(days=1, hours=4), "exit", 2),
            self._punch("100", self._at(days=2, hours=0, minutes=10), "entry", 3),
            self._punch("100", self._at(days=2, hours=4), "exit", 4),
        ])
        summary = self._book(self._report(people, self.day, self.day + timedelta(days=1)))["Person Summary"]
        values = [c.value for c in summary[4]]
        self.assertEqual(values[4].strftime("%H:%M"), "00:00")
        self.assertEqual(values[3], 2)  # days present

    def test_csv_is_daily_detail_with_bom_and_hh_mm(self) -> None:
        people = self._people(1)
        self.database.insert_punches([
            self._punch("100", self._at(hours=22, minutes=10, seconds=5), "entry", 1),
            self._punch("100", self._at(days=1, hours=3, minutes=40, seconds=50), "exit", 2),
        ])
        text = export.to_csv(self._report(people, self.day, self.day))
        self.assertTrue(text.startswith("\ufeff"))
        rows = list(csv.reader(io.StringIO(text[1:])))
        self.assertEqual(rows[0], list(export.DAILY_HEADER))
        self.assertEqual(len(rows), 2)
        row = dict(zip(rows[0], rows[1]))
        self.assertEqual((row["First In"], row["Last Out"]), ("22:10", "03:40"))
        self.assertEqual(row["Effective"], "5:30")
        self.assertEqual(row["Effective (h)"], "5.51")
        self.assertEqual(row["Date"], "2026-09-01")

    def test_formula_text_is_prefixed_in_csv_and_every_sheet(self) -> None:
        self.database.upsert_employee(device_user_id="7", name='=HYPERLINK("x")', email="h@x.com")
        report = self._report(self.database.employees(), self.day, self.day)
        parsed = list(csv.DictReader(io.StringIO(export.to_csv(report)[1:])))
        self.assertEqual(parsed[0]["Name"], '\'=HYPERLINK("x")')
        book = self._book(report)
        for sheet in book.worksheets:
            cell = sheet.cell(row=4, column=2)
            self.assertEqual(cell.value, '\'=HYPERLINK("x")', sheet.title)
            self.assertEqual(cell.data_type, "s")
        for prefix in ("+", "-", "@", "\t", "\r"):
            self.assertEqual(export.safe_text(prefix + "a"), "'" + prefix + "a")
        self.assertEqual(export.safe_text("plain"), "plain")

    def test_range_query_reaches_into_the_next_morning(self) -> None:
        people = self._people(1)
        self.database.insert_punches([
            self._punch("100", self._at(hours=22), "entry", 1),
            self._punch("100", self._at(days=1, hours=5, minutes=59), "exit", 2),
            self._punch("100", self._at(hours=5), "entry", 3),  # previous work day
        ])
        (row,) = self._report(people, self.day, self.day).rows
        self.assertEqual(row.effective, timedelta(hours=7, minutes=59))

    def test_day_start_hour_zero_restores_calendar_days(self) -> None:
        people = self._people(1)
        self.database.insert_punches([
            self._punch("100", self._at(hours=22, minutes=10), "entry", 1),
            self._punch("100", self._at(days=1, hours=3, minutes=40), "exit", 2),
        ])
        report = self._report(people, self.day, self.day + timedelta(days=1), day_start_hour=0)
        self.assertEqual([r.code for r in report.rows], ["I", "I"])
        self.assertIn("00:00 to 00:00", report.note)

    def test_range_validation(self) -> None:
        today = date(2026, 10, 10)
        self.assertEqual(export.parse_range(None, None, today), (today, today))
        self.assertEqual(export.parse_range("2026-07-01", "2026-10-01", today)[1], date(2026, 10, 1))  # 93 days
        for bad in (("2026-10-02", "2026-10-01"), ("2026-06-30", "2026-10-01"), ("nope", "2026-10-01")):
            with self.assertRaises(export.RangeError):
                export.parse_range(*bad, today)

    def test_each_punch_uses_its_devices_quick_swipe_window(self) -> None:
        # KTD3: 09:00:45 is collapsed at 60 s and counts as Out at the X990's 30 s floor.
        people = self._people(1)
        self.database.insert_punches([
            self._punch("100", self._at(hours=9, seconds=s), None, n)
            for n, s in enumerate((0, 45, 3 * 3600, 8 * 3600))
        ])

        def effective() -> int:
            results = export.day_results(self.database, people, self.day, self.day, target_minutes=480,
                                         today=date(2026, 10, 10), now=None)
            return results[(people[0]["id"], self.day)].effective_seconds

        self.assertEqual(effective(), 3 * 3600)
        self.database.set_device_policy(DEVICES[0].id, window_seconds=30, wait_for_door_close=True)
        self.assertEqual(effective(), 45 + 5 * 3600)

    def test_query_count_is_constant_for_10_vs_100_people(self) -> None:
        statements: list[str] = []
        original = db_module.Database._open

        def counting_open(database):
            connection = original(database)
            connection.set_trace_callback(
                lambda sql: statements.append(sql) if not sql.lstrip().upper().startswith("PRAGMA") else None
            )
            return connection

        def run(people, days) -> int:
            statements.clear()
            with mock.patch.object(db_module.Database, "_open", counting_open):
                report = self._report(people, self.day, self.day + timedelta(days=days - 1))
                export.to_xlsx(report)
                export.to_csv(report)
            self.assertEqual(len(report.rows), len(people) * days)
            return len(statements)

        everyone = self._people(100)
        self.database.insert_punches(
            [self._punch(p["device_user_id"], self._at(days=d, hours=9), "entry", d) for p in everyone for d in range(31)]
        )
        small = run(everyone[:10], 31)
        large = run(everyone, 31)
        self.assertEqual(small, large)
        self.assertLessEqual(large, 1)


    def test_approved_fixes_change_the_report_and_group_under_the_owner(self) -> None:
        # KTD4: Person 000 swipes on her extra ID; HR adds an In, re-directs a
        # wrong Out and ignores a duplicate. The device rows never change.
        (person,) = self._people(1)
        self.database.add_extra_device_id(person["id"], "NP0100")
        self.database.insert_punches([
            self._punch("NP0100", self._at(hours=9), "entry", 1),
            self._punch("NP0100", self._at(hours=13), "exit", 2),
            self._punch("NP0100", self._at(hours=14), "exit", 3),
            self._punch("NP0100", self._at(hours=14, minutes=30), "exit", 4),
            self._punch("NP0100", self._at(hours=18), "exit", 5),
        ])
        (person,) = self.database.employees()
        before = self._report([person], self.day, self.day).rows[0]
        self.assertEqual((before.code, before.effective), ("R", timedelta(hours=4)))
        request_id = self.database.file_fix_lines(person["id"], self.day, [
            {"kind": "flip", "target_fingerprint": "fp-NP0100-3", "direction": "entry"},
            {"kind": "ignore", "target_fingerprint": "fp-NP0100-4"},
            {"kind": "add_in", "event_time": self._at(hours=19)},
            {"kind": "add_out", "event_time": self._at(hours=20)},
        ], source="session")["request_id"]
        self.database.upsert_employee(device_user_id="900", name="HR", email="hr@x.com")
        hr = self.database.employee_by_device_user_id("900")
        self.database.set_employee_role(hr["id"], "hr")
        self.database.decide_fix_lines(request_id, [
            {"line_id": l["id"], "status": "approved"}
            for l in self.database.fix_request(request_id)["lines"]], actor_id=hr["id"])
        after = self._report([person], self.day, self.day).rows[0]
        self.assertEqual((after.code, after.status), ("P", "Met"))
        self.assertEqual(after.effective, timedelta(hours=9))
        self.assertEqual((after.first_in, after.last_out), (self._at(hours=9), self._at(hours=20)))
        self.assertEqual(after.punches, 6)  # the ignored duplicate does not count
        self.assertEqual(after.sites, ("s",))

if __name__ == "__main__":
    unittest.main()
